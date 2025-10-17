from typing import List, Dict, Any
import asyncio
from io import BytesIO
import PyPDF2
import docx
import fitz  # PyMuPDF for better PDF handling
import uuid
import re
import numpy as np
from datetime import datetime
from fastapi import UploadFile
import os

import logging

logger = logging.getLogger(__name__)

# Simple storage for development (replace with database in production)
document_storage = {}

class DocumentProcessor:
    def __init__(self):
        self.supported_formats = {
            'application/pdf': self._extract_pdf_text,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document': self._extract_docx_text,
            'application/msword': self._extract_doc_text,
            'text/plain': self._extract_text_text
        }
        logger.info("Document processor initialized with semantic chunking capabilities")

    async def process_file(self, file: UploadFile) -> Dict[str, Any]:
        """Process uploaded file and store for further processing"""
        logger.info(f"Processing uploaded file: {file.filename}")
        try:
            # Generate unique document ID
            document_id = f"doc_{uuid.uuid4().hex[:8]}"
            # Read file content
            file_content = await file.read()
            logger.info(f"Read {len(file_content)} bytes from file")
            # Extract text based on file type
            content_type = file.content_type or self._detect_content_type(file.filename)
            logger.info(f"Detected content type: {content_type}")
            extracted_text = self.extract_text(file_content, content_type)
            logger.info(f"Extracted {len(extracted_text)} characters of text")
            # Store document data
            document_data = {
                'document_id': document_id,
                'filename': file.filename,
                'content_type': content_type,
                'raw_content': extracted_text,
                'file_size': len(file_content),
                'created_at': datetime.now().isoformat(),
                'status': 'uploaded'
            }
            document_storage[document_id] = document_data
            logger.info(f"Stored document data for: {document_id}")
            return document_data
        except Exception as e:
            logger.error(f"File processing failed: {e}")
            raise Exception(f"File processing failed: {e}")
    
    async def get_document_content(self, document_id: str) -> str:
        """Get document content by ID"""
        logger.info(f"Retrieving content for document: {document_id}")
        if document_id not in document_storage:
            raise Exception(f"Document {document_id} not found")
        content = document_storage[document_id]['raw_content']
        logger.info(f"Retrieved {len(content)} characters for document {document_id}")
        return content
    
    async def semantic_chunk(self, content: str, document_id: str) -> List[Dict[str, Any]]:
        """Chunk document content into semantically coherent pieces - OPTIMIZED for fewer API calls"""
        logger.info(f"Starting OPTIMIZED semantic chunking for document: {document_id}")
        logger.info(f"Input content length: {len(content)} characters")
        try:
            chunks = []
            # Method 1: Split by sections/headers
            section_chunks = self._split_by_sections(content)
            logger.info(f"Section-based chunking created {len(section_chunks)} chunks")
            
            # Method 2: OPTIMIZATION - Create larger chunks (3000-4000 chars) to reduce API calls
            refined_chunks = []
            for i, section_chunk in enumerate(section_chunks):
                if len(section_chunk) > 4000:  # Only split very large sections
                    sentence_chunks = self._split_by_sentences(section_chunk, max_length=3500)
                    refined_chunks.extend(sentence_chunks)
                    logger.debug(f"Split large section {i} into {len(sentence_chunks)} sentence chunks")
                else:
                    refined_chunks.append(section_chunk)
            
            # Method 3: OPTIMIZATION - Combine small chunks to reduce total number
            combined_chunks = []
            current_chunk = ""
            
            for chunk in refined_chunks:
                # If adding this chunk would still be under our target size, combine it
                if len(current_chunk) + len(chunk) < 3000:
                    current_chunk += "\n\n" + chunk if current_chunk else chunk
                else:
                    # Save current chunk and start new one
                    if current_chunk.strip():
                        combined_chunks.append(current_chunk.strip())
                    current_chunk = chunk
            
            # Don't forget the last chunk
            if current_chunk.strip():
                combined_chunks.append(current_chunk.strip())
            
            logger.info(f"OPTIMIZED chunking: {len(section_chunks)} -> {len(refined_chunks)} -> {len(combined_chunks)} chunks")
            
            # Create chunk objects with metadata
            for i, chunk_text in enumerate(combined_chunks):
                if chunk_text.strip():  # Only non-empty chunks
                    # --- ADDED: Robust Source Identifier ---
                    # The source identifier should be contextually meaningful (e.g., section, US ID)
                    # Here we mock it, but the actual implementation would need deep parsing.
                    # We rely on the chunk index and document ID.
                    source_identifier = f"{document_id}_src_{i:03d}"
                    
                    chunk_data = {
                        'chunk_id': f"{document_id}_chunk_{i:03d}",
                        'content': chunk_text.strip(),
                        'text': chunk_text.strip(),  # Alias for compatibility
                        'document_id': document_id,
                        'chunk_index': i,
                        'source_identifier': source_identifier,  # New field to track source
                        'metadata': {
                            'length': len(chunk_text.strip()),
                            'word_count': len(chunk_text.strip().split()),
                            'sentence_count': len([s for s in chunk_text.split('.') if s.strip()]),
                            'type': self._classify_chunk_type(chunk_text.strip()),
                            'source_section': f"Section {i//5 + 1}"  # Example heuristic
                        }
                    }
                    chunks.append(chunk_data)
            
            # Log chunking statistics
            if chunks:
                avg_length = sum(len(chunk['content']) for chunk in chunks) / len(chunks)
                max_length = max(len(chunk['content']) for chunk in chunks)
                min_length = min(len(chunk['content']) for chunk in chunks)
                logger.info(f"OPTIMIZED Chunking Statistics:")
                logger.info(f"   - Total chunks: {len(chunks)} (REDUCED from original)")
                logger.info(f"   - Average length: {avg_length:.1f} characters (INCREASED for efficiency)")
                logger.info(f"   - Max length: {max_length} characters")
                logger.info(f"   - Min length: {min_length} characters")
                logger.info(f"   - Estimated API calls needed: {len(chunks)} (DOWN from {len(refined_chunks)})")
                
                # Show sample chunks
                for i, chunk in enumerate(chunks[:3]):
                    preview = chunk['content'][:100] + "..." if len(chunk['content']) > 100 else chunk['content']
                    logger.debug(f"   - Chunk {i}: {preview}")
            
            logger.info(f"OPTIMIZED semantic chunking completed: {len(chunks)} chunks created")
            return chunks
            return chunks
        except Exception as e:
            logger.error(f"Semantic chunking failed: {e}")
            raise Exception(f"Semantic chunking failed: {e}")
    
    def _split_by_sections(self, content: str) -> List[str]:
        """Split content by logical sections (headers, paragraphs, etc.)"""
        
        # Look for section headers (various formats)
        section_patterns = [
            r'\n\s*(?:\d+\.?\s+)?[A-Z][^.!?]*:?\s*\n',  # Numbered or titled sections
            r'\n\s*(?:SECTION|Chapter|Part)\s+\d+[^\n]*\n',  # Explicit sections
            r'\n\s*[A-Z][A-Z\s]{3,}[A-Z]\s*\n',  # ALL CAPS headers
            r'\n\s*\d+\.\d+[^\n]*\n',  # Numbered subsections
            r'\n\s*[IVX]+\.\s+[^\n]*\n'  # Roman numeral sections
        ]
        
        # Try to find section breaks
        sections = []
        current_section = ""
        
        lines = content.split('\n')
        
        for line in lines:
            # Check if this line is a section header
            is_header = False
            for pattern in section_patterns:
                if re.match(pattern, '\n' + line + '\n'):
                    is_header = True
                    break
            
            if is_header and current_section.strip():
                # Start new section
                sections.append(current_section.strip())
                current_section = line + '\n'
            else:
                current_section += line + '\n'
        
        # Add final section
        if current_section.strip():
            sections.append(current_section.strip())
        
        # If no sections found, split by double newlines (paragraphs)
        if len(sections) <= 1:
            paragraph_sections = [p.strip() for p in content.split('\n\n') if p.strip()]
            sections = paragraph_sections if len(paragraph_sections) > 1 else [content]
        
        return sections
    
    def _split_by_sentences(self, text: str, max_length: int = 1500) -> List[str]:
        """Split text by sentences while respecting max length"""
        
        # Simple sentence splitting (could be enhanced with NLTK)
        sentences = re.split(r'(?<=[.!?])\s+', text)
        
        chunks = []
        current_chunk = ""
        
        for sentence in sentences:
            # Check if adding this sentence would exceed max length
            if len(current_chunk) + len(sentence) > max_length and current_chunk:
                chunks.append(current_chunk.strip())
                current_chunk = sentence
            else:
                current_chunk += (" " if current_chunk else "") + sentence
        
        # Add final chunk
        if current_chunk.strip():
            chunks.append(current_chunk.strip())
        
        return chunks
    
    def _classify_chunk_type(self, text: str) -> str:
        """Classify chunk type based on content patterns"""
        text_lower = text.lower()
        
        # Simple classification rules
        if any(keyword in text_lower for keyword in ['requirement', 'shall', 'must', 'should']):
            return 'requirement'
        elif any(keyword in text_lower for keyword in ['test', 'verify', 'validate', 'check']):
            return 'test_related'
        elif any(keyword in text_lower for keyword in ['user', 'customer', 'stakeholder']):
            return 'user_story'
        elif any(keyword in text_lower for keyword in ['system', 'architecture', 'design']):
            return 'technical'
        elif any(keyword in text_lower for keyword in ['business', 'process', 'workflow']):
            return 'business_logic'
        else:
            return 'general'
    
    def _detect_content_type(self, filename: str) -> str:
        """Detect content type from filename"""
        if not filename:
            return 'text/plain'
        
        extension = filename.lower().split('.')[-1] if '.' in filename else ''
        
        content_type_map = {
            'pdf': 'application/pdf',
            'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            'doc': 'application/msword',
            'txt': 'text/plain'
        }
        
        return content_type_map.get(extension, 'text/plain')

    def _extract_pdf_text(self, file_content: bytes) -> str:
        """Extract text from PDF using PyMuPDF (better than PyPDF2)"""
        try:
            pdf_document = fitz.open(stream=file_content, filetype="pdf")
            text = ""
            
            for page_num in range(pdf_document.page_count):
                page = pdf_document[page_num]
                text += page.get_text()
                text += "\n\n"  # Add page break
            
            pdf_document.close()
            return text.strip()
            
        except Exception as e:
            logger.error(f"Error extracting PDF text with PyMuPDF: {str(e)}")
            # Fallback to PyPDF2
            try:
                pdf_file = BytesIO(file_content)
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                text = ""
                
                for page in pdf_reader.pages:
                    text += page.extract_text()
                    text += "\n\n"
                
                return text.strip()
            except Exception as e2:
                logger.error(f"Error extracting PDF text with PyPDF2: {str(e2)}")
                raise Exception("Could not extract text from PDF file")

    def _extract_docx_text(self, file_content: bytes) -> str:
        """Extract text from DOCX file"""
        try:
            doc_file = BytesIO(file_content)
            doc = docx.Document(doc_file)
            text = ""
            
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
            
            return text.strip()
            
        except Exception as e:
            logger.error(f"Error extracting DOCX text: {str(e)}")
            raise Exception("Could not extract text from DOCX file")

    def _extract_doc_text(self, file_content: bytes) -> str:
        """Extract text from legacy DOC file"""
        # Note: This is a simplified implementation
        # For full DOC support, you'd need python-docx2txt or antiword
        try:
            # For now, just return an error message
            raise Exception("Legacy DOC files are not fully supported. Please convert to DOCX or PDF.")
        except Exception as e:
            logger.error(f"Error extracting DOC text: {str(e)}")
            raise e

    def _extract_text_text(self, file_content: bytes) -> str:
        """Extract text from plain text file"""
        try:
            # Try different encodings
            encodings = ['utf-8', 'utf-16', 'latin-1', 'cp1252']
            
            for encoding in encodings:
                try:
                    return file_content.decode(encoding)
                except UnicodeDecodeError:
                    continue
            
            # If all encodings fail, use utf-8 with error handling
            return file_content.decode('utf-8', errors='replace')
            
        except Exception as e:
            logger.error(f"Error extracting text: {str(e)}")
            raise Exception("Could not extract text from file")

    def extract_text(self, file_content: bytes, content_type: str) -> str:
        """Extract text from file based on content type"""
        try:
            if content_type not in self.supported_formats:
                raise Exception(f"Unsupported file format: {content_type}")
            
            extractor = self.supported_formats[content_type]
            text = extractor(file_content)
            
            if not text or len(text.strip()) < 10:
                raise Exception("No meaningful text content found in file")
            
            logger.info(f"Extracted {len(text)} characters from {content_type} file")
            return text
            
        except Exception as e:
            logger.error(f"Error extracting text: {str(e)}")
            raise e

    async def process_document(self, document_id: str, content: str, progress_callback=None) -> Dict[str, Any]:
        """Process document and generate test recommendations using clustering and rule-based logic only."""
        try:
            logger.info(f"[START] Starting comprehensive document processing for: {document_id}")
            logger.info(f"[CONTENT] Document content length: {len(content)} characters")
            
            # Step 1: Chunk the text
            from .embedding_service import EmbeddingService
            from .clustering_service import ClusteringService
            embedding_service = EmbeddingService()
            clustering_service = ClusteringService()
            
            if progress_callback:
                await progress_callback(10, "Starting text analysis...")

            logger.info("[STEP 1] Performing semantic chunking...")
            # Use semantic_chunk for chunking
            chunks = await self.semantic_chunk(content, document_id)
            logger.info(f"[SUCCESS] Document chunked into {len(chunks)} pieces")
            logger.info(f"[DEBUG] Chunk size distribution:")
            chunk_sizes = [len(chunk['content']) for chunk in chunks]
            if chunk_sizes:
                logger.info(f"   - Average chunk size: {np.mean(chunk_sizes):.0f} chars")
                logger.info(f"   - Min chunk size: {min(chunk_sizes)} chars")
                logger.info(f"   - Max chunk size: {max(chunk_sizes)} chars")
            
            if progress_callback:
                await progress_callback(20, "Text chunking completed")

            # Step 2: Generate embeddings
            logger.info("[STEP 2] Generating embeddings for chunks...")
            embeddings = await embedding_service.generate_embeddings(chunks)
            logger.info(f"[SUCCESS] Generated {len(embeddings)} embeddings")
            
            if progress_callback:
                await progress_callback(40, "Generated embeddings")

            # Step 3: Cluster embeddings
            logger.info("[STEP 3] Clustering embeddings into feature groups...")
            clusters = await clustering_service.cluster_embeddings(embeddings, chunks)
            n_clusters = len(clusters)
            logger.info(f"[SUCCESS] Created {n_clusters} feature clusters")
            
            # Log detailed cluster information
            logger.info("[CLUSTERING] Cluster Analysis Summary:")
            for i, cluster in enumerate(clusters):
                cluster_id = cluster.get('cluster_id', i)
                size = cluster.get('size', 0)
                representative = cluster.get('representative_text', '')[:100]
                logger.info(f"   -> Cluster {cluster_id}: {size} chunks - \"{representative}...\"")
            
            if progress_callback:
                await progress_callback(60, f"Created {n_clusters} feature clusters")

            # Step 4: Domain-aware test recommendations using the main engine
            logger.info("[STEP 4] Generating domain-aware test recommendations...")
            from .domain_aware_test_engine import DomainAwareTestEngine
            test_engine = DomainAwareTestEngine()
            
            # Generate domain-aware test recommendations for all clusters
            logger.info("[TEST-GEN] Running domain-aware test analysis on clusters...")
            cluster_analyses = test_engine.analyze_clusters(clusters)
            standard_tests, recommended_tests = test_engine.generate_recommendations(cluster_analyses)
            test_results = test_engine.format_recommendations(standard_tests, recommended_tests)
            
            total_recommendations = len(test_results.get('recommendations', []))
            logger.info(f"[SUCCESS] Domain-aware test engine completed")
            logger.info(f"[SUMMARY] Generated {total_recommendations} cluster recommendations")
            
            # Convert to the expected format for compatibility
            logger.info("[STEP 5] Formatting results for API response...")
            
            # The domain-aware test engine returns recommendations in the expected format
            recommendations = test_results.get('recommendations', [])
            logger.info(f"[FORMAT] Using domain-aware recommendations: {len(recommendations)} clusters")
            
            if progress_callback:
                await progress_callback(80, f"Formatted {len(recommendations)} cluster recommendations")

            # Final result compilation
            total_tests = sum(len(rec.get("tests", [])) for rec in recommendations)
            logger.info(f"[SUCCESS] Document processing completed successfully!")
            logger.info(f"[SUMMARY] Final Results Summary:")
            logger.info(f"   -> Clusters generated: {len(recommendations)}")
            logger.info(f"   -> Total test recommendations: {total_tests}")
            logger.info(f"   -> Processing method: Rule-based clustering + domain-aware test engine")
            
            result = {
                "document_id": document_id,
                "status": "completed",
                "clusters_generated": len(recommendations),
                "total_tests": total_tests,
                "recommendations": recommendations,
                "comprehensive_analysis": {
                    "total_test_types_identified": total_tests,
                    "engine_metadata": {
                        "engine_type": "domain_aware_test_engine",
                        "clusters_analyzed": len(cluster_analyses),
                        "processing_method": "semantic_clustering_with_domain_awareness"
                    }
                }
            }
            
            return result
            
        except Exception as e:
            logger.error(f"[ERROR] Error processing document {document_id}: {str(e)}")
            if progress_callback:
                await progress_callback(-1, f"Processing failed: {str(e)}")
            raise e

# Global instance
document_processor = DocumentProcessor()